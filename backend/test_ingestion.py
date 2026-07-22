"""
Integration test script for document ingestion and vector store operations.
Creates dummy PDF, DOCX, and TXT files, chunks them, inserts into ChromaDB, and queries.
"""

import os
from pathlib import Path
import docx
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter

from core.ingestion import extract_text, chunk_document
from core.vector_store import VectorStoreManager

def create_sample_files():
    print("Creating sample files for testing...")
    
    # 1. Create a dummy TXT file
    txt_path = Path("sample.txt")
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(
            "Agentic RAG Research Assistant Project Overview.\n"
            "This project implements a multi-agent RAG pipeline using LangGraph.\n"
            "The system is designed to retrieve document chunks and grade their relevance.\n"
            "It also checks for hallucinations in the final answer."
        )
    print(f"Created {txt_path.absolute()}")

    # 2. Create a dummy DOCX file
    docx_path = Path("sample.docx")
    doc = docx.Document()
    doc.add_heading("Agentic RAG Design Document", level=0)
    doc.add_paragraph(
        "The architecture consists of a FastAPI backend and a Streamlit frontend. "
        "ChromaDB is used as the vector database because it runs locally and is easy to set up."
    )
    doc.add_paragraph(
        "A key component is the Relevance Grader agent, which evaluates whether a retrieved chunk "
        "is relevant to the user's question. This prevents the generator from using noisy information."
    )
    doc.save(docx_path)
    print(f"Created {docx_path.absolute()}")
    
    # 3. Create a dummy multi-page PDF file
    pdf_path = Path("sample.pdf")
    c = canvas.Canvas(str(pdf_path), pagesize=letter)
    
    # Page 1 content
    c.drawString(100, 750, "Agentic RAG Research Assistant PDF Document - Page 1")
    c.drawString(100, 700, "This is page one text. We talk about Query Rewriter.")
    c.drawString(100, 680, "The Query Rewriter agent reformulates raw input search queries.")
    c.showPage()
    
    # Page 2 content
    c.drawString(100, 750, "Agentic RAG Research Assistant PDF Document - Page 2")
    c.drawString(100, 700, "This is page two text. We talk about Hallucination Checker.")
    c.drawString(100, 680, "The Hallucination Checker validates generated answers against retrieved sources.")
    c.showPage()
    
    c.save()
    print(f"Created {pdf_path.absolute()}")
    
    return txt_path, docx_path, pdf_path

def clean_sample_files(files):
    for f in files:
        if os.path.exists(f):
            os.remove(f)
            print(f"Removed {f}")

def main():
    txt_file, docx_file, pdf_file = create_sample_files()
    
    try:
        # Ingest TXT
        print("\n--- Testing TXT Ingestion ---")
        txt_pages = extract_text(txt_file)
        print("TXT extracted pages:", len(txt_pages))
        txt_chunks = chunk_document(txt_pages, doc_id="doc_txt_01", doc_name="sample.txt")
        print(f"TXT chunks created: {len(txt_chunks)}")
        for idx, chunk in enumerate(txt_chunks):
            print(f"Chunk {idx}: {chunk['chunk_id']} (Page {chunk['page_number']}): {chunk['chunk_text'][:60]}...")

        # Ingest DOCX
        print("\n--- Testing DOCX Ingestion ---")
        docx_pages = extract_text(docx_file)
        print("DOCX extracted pages:", len(docx_pages))
        docx_chunks = chunk_document(docx_pages, doc_id="doc_docx_01", doc_name="sample.docx")
        print(f"DOCX chunks created: {len(docx_chunks)}")
        for idx, chunk in enumerate(docx_chunks):
            print(f"Chunk {idx}: {chunk['chunk_id']} (Page {chunk['page_number']}): {chunk['chunk_text'][:60]}...")

        # Ingest PDF
        print("\n--- Testing PDF Ingestion ---")
        pdf_pages = extract_text(pdf_file)
        print("PDF extracted pages:", len(pdf_pages))
        # Ensure we have 2 pages
        assert len(pdf_pages) == 2, f"Expected 2 pages, got {len(pdf_pages)}"
        print("Page 1 preview:", pdf_pages[0]['text'][:60].strip())
        print("Page 2 preview:", pdf_pages[1]['text'][:60].strip())
        
        pdf_chunks = chunk_document(pdf_pages, doc_id="doc_pdf_01", doc_name="sample.pdf")
        print(f"PDF chunks created: {len(pdf_chunks)}")
        for idx, chunk in enumerate(pdf_chunks):
            print(f"Chunk {idx}: {chunk['chunk_id']} (Page {chunk['page_number']}): {chunk['chunk_text'][:60]}...")

        # Test Vector Store
        print("\n--- Testing Vector Store Manager ---")
        manager = VectorStoreManager()
        
        # Clean existing test collection if any
        try:
            manager.delete_by_doc_id("doc_txt_01")
            manager.delete_by_doc_id("doc_docx_01")
            manager.delete_by_doc_id("doc_pdf_01")
        except Exception:
            pass
            
        print("Adding chunks to ChromaDB...")
        manager.add_chunks(txt_chunks)
        manager.add_chunks(docx_chunks)
        manager.add_chunks(pdf_chunks)
        print("Chunks added successfully.")
        
        # Query for Query Rewriter (should match Page 1 of PDF)
        query_text = "How does the Query Rewriter behave?"
        print(f"\nQuerying: '{query_text}'")
        results = manager.query(query_text, top_k=2)
        print(f"Found {len(results)} results:")
        for idx, res in enumerate(results):
            print(f"Result {idx+1}:")
            print(f"  Doc ID: {res['doc_id']}")
            print(f"  Doc Name: {res['doc_name']}")
            print(f"  Page: {res['page_number']}")
            print(f"  Text: {res['chunk_text']}")
            print(f"  Score: {res['score']}")
            
        # Query for Hallucination Checker (should match Page 2 of PDF)
        query_text_2 = "What does the Hallucination Checker do?"
        print(f"\nQuerying: '{query_text_2}'")
        results_2 = manager.query(query_text_2, top_k=2)
        print(f"Found {len(results_2)} results:")
        for idx, res in enumerate(results_2):
            print(f"Result {idx+1}:")
            print(f"  Doc ID: {res['doc_id']}")
            print(f"  Doc Name: {res['doc_name']}")
            print(f"  Page: {res['page_number']}")
            print(f"  Text: {res['chunk_text']}")

    finally:
        clean_sample_files([txt_file, docx_file, pdf_file])

if __name__ == "__main__":
    main()
